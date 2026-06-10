from __future__ import annotations

from pathlib import Path
from textwrap import wrap

from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
ASSETS = REPORTS / "presentation_assets"
VISUAL_DIR = ASSETS / "visual_slides"
PPTX_PATH = REPORTS / "gangwon_emergency_hospital_guide_presentation.pptx"
NOTES_PATH = REPORTS / "presentation_speaker_notes.md"

W, H = 1920, 1080

FONT = Path("C:/Windows/Fonts/malgun.ttf")
FONT_BOLD = Path("C:/Windows/Fonts/malgunbd.ttf")


def f(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(FONT_BOLD if bold else FONT), size)


NAVY = "#08142c"
NAVY2 = "#0f2344"
BLUE = "#1c6cff"
CYAN = "#00cbe8"
MINT = "#20d6a1"
RED = "#ff4d5a"
YELLOW = "#ffc247"
WHITE = "#ffffff"
BG = "#f5f8fb"
PANEL = "#ffffff"
INK = "#16233a"
MUTED = "#6b7c93"
LINE = "#d9e5ee"


def canvas(bg=BG):
    return Image.new("RGB", (W, H), bg)


def draw_text(draw, xy, text, size=42, color=INK, bold=False, anchor=None, align="left", spacing=8):
    draw.multiline_text(xy, text, font=f(size, bold), fill=color, anchor=anchor, align=align, spacing=spacing)


def text_box(draw, xy, wh, text, size=34, color=INK, bold=False, align="left", line_chars=18, spacing=8):
    x, y = xy
    lines = []
    for raw in text.split("\n"):
        if not raw:
            lines.append("")
        else:
            lines.extend(wrap(raw, width=line_chars, break_long_words=False))
    draw.multiline_text((x, y), "\n".join(lines), font=f(size, bold), fill=color, align=align, spacing=spacing)


def rounded(draw, box, r=34, fill=PANEL, outline=None, width=2):
    draw.rounded_rectangle(box, radius=r, fill=fill, outline=outline, width=width)


def shadowed_card(img, box, r=34, fill=PANEL, outline=LINE):
    shadow = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    sd = ImageDraw.Draw(shadow)
    shifted = (box[0] + 10, box[1] + 14, box[2] + 10, box[3] + 14)
    sd.rounded_rectangle(shifted, radius=r, fill=(20, 40, 80, 34))
    shadow = shadow.filter(ImageFilter.GaussianBlur(18))
    img.paste(Image.alpha_composite(img.convert("RGBA"), shadow).convert("RGB"))
    d = ImageDraw.Draw(img)
    rounded(d, box, r, fill, outline, 2)


def pill(draw, box, text, fill=CYAN, color=NAVY, size=24):
    rounded(draw, box, r=(box[3] - box[1]) // 2, fill=fill, outline=fill)
    draw.text(((box[0] + box[2]) / 2, (box[1] + box[3]) / 2 - 2), text, font=f(size, True), fill=color, anchor="mm")


def title(draw, main, sub=None, color=INK):
    draw_text(draw, (110, 78), main, 56, color, True)
    if sub:
        draw_text(draw, (113, 154), sub, 25, MUTED, False)


def footer(draw, n):
    draw.text((1810, 1015), f"{n:02d}", font=f(20, True), fill="#a8b4c5", anchor="ra")


def arrow(draw, start, end, color=CYAN, width=8):
    draw.line([start, end], fill=color, width=width)
    x1, y1 = start
    x2, y2 = end
    if x2 >= x1:
        pts = [(x2, y2), (x2 - 28, y2 - 16), (x2 - 28, y2 + 16)]
    else:
        pts = [(x2, y2), (x2 + 28, y2 - 16), (x2 + 28, y2 + 16)]
    draw.polygon(pts, fill=color)


def paste_fit(img, path, box):
    src = Image.open(path).convert("RGB")
    x1, y1, x2, y2 = box
    bw, bh = x2 - x1, y2 - y1
    src.thumbnail((bw, bh), Image.LANCZOS)
    x = x1 + (bw - src.width) // 2
    y = y1 + (bh - src.height) // 2
    img.paste(src, (x, y))


def add_metric(draw, box, value, label, accent=CYAN):
    rounded(draw, box, 32, WHITE, LINE)
    x1, y1, x2, y2 = box
    draw.text(((x1 + x2) / 2, y1 + 68), value, font=f(58, True), fill=accent, anchor="mm")
    draw.text(((x1 + x2) / 2, y1 + 138), label, font=f(24, True), fill=MUTED, anchor="mm")


def slide_01():
    img = canvas(NAVY)
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, W, H), fill=NAVY)
    d.ellipse((1280, -240, 2180, 660), fill="#12345f")
    d.ellipse((1370, -140, 2070, 560), fill="#0bbbd833")
    pill(d, (110, 94, 430, 150), "DATA MINING PROJECT", CYAN, NAVY, 25)
    draw_text(d, (112, 255), "강원도 맞춤형\n응급 및 상시 병원\n안내 시스템", 74, WHITE, True, spacing=18)
    draw_text(d, (118, 610), "자연어 증상 입력에서 응급도와 병원 추천까지", 34, "#c9e8f0", False)
    shadowed_card(img, (1050, 230, 1760, 795), 44, NAVY2, "#24466d")
    paste_fit(img, ASSETS / "site_home_top.png", (1090, 270, 1720, 755))
    footer(d, 1)
    return img


def slide_02():
    img = canvas()
    d = ImageDraw.Draw(img)
    title(d, "문제 정의", "증상은 자연어이고, 병원 선택은 복잡하다")
    shadowed_card(img, (110, 285, 510, 675), 36)
    shadowed_card(img, (760, 235, 1160, 725), 36, "#eafcff", CYAN)
    shadowed_card(img, (1410, 285, 1810, 675), 36)
    draw_text(d, (310, 375), "사용자", 42, INK, True, anchor="mm")
    draw_text(d, (310, 505), "“말이 어눌하고\n팔에 힘이 없어요”", 30, MUTED, False, anchor="mm", align="center")
    draw_text(d, (960, 330), "비어 있던\n중간 판단", 48, NAVY, True, anchor="mm", align="center")
    draw_text(d, (960, 515), "증상군 · 진료과\n의심 질환 · 응급도", 33, BLUE, True, anchor="mm", align="center")
    draw_text(d, (1610, 390), "결과", 42, INK, True, anchor="mm")
    draw_text(d, (1610, 515), "어디로\n가야 하는가", 33, MUTED, False, anchor="mm", align="center")
    arrow(d, (525, 480), (735, 480), CYAN)
    arrow(d, (1185, 480), (1385, 480), CYAN)
    draw_text(d, (960, 865), "목표: 증상과 병원 사이의 판단 과정을 데이터마이닝으로 구현", 40, NAVY, True, anchor="mm")
    footer(d, 2)
    return img


def slide_03():
    img = canvas()
    d = ImageDraw.Draw(img)
    title(d, "프로젝트 목표", "많이 묻지 않고, 위험한 증상은 놓치지 않는다")
    goals = [
        ("01", "모델 기반 예측", "증상군·진료과·질환 후보"),
        ("02", "최소 문진", "필요한 질문만 1~4개"),
        ("03", "응급도 판단", "red flag 우선 보정"),
        ("04", "병원 추천", "병상·거리·ETA 반영"),
    ]
    for i, (num, head, body) in enumerate(goals):
        x = 130 + i * 445
        shadowed_card(img, (x, 295, x + 360, 655), 38)
        d.text((x + 42, 350), num, font=f(42, True), fill=CYAN)
        d.text((x + 42, 438), head, font=f(36, True), fill=NAVY)
        text_box(d, (x + 45, 525), (280, 100), body, 27, MUTED, False, line_chars=12)
    rounded(d, (255, 800, 1665, 910), 34, NAVY, NAVY)
    d.text((960, 855), "질환 판단은 모델 중심 · 응급도 데이터는 위험도 계산 보조", font=f(39, True), fill=WHITE, anchor="mm")
    footer(d, 3)
    return img


def slide_04():
    img = canvas()
    d = ImageDraw.Draw(img)
    title(d, "전체 시스템 구조", "입력부터 결과까지 하나의 파이프라인")
    labels = [
        ("증상 입력", "자연어 문장"),
        ("모델 예측", "질환 TOP3"),
        ("문진 선택", "질환별 질문"),
        ("응급도 계산", "red flag + risk"),
        ("병원 추천", "진료과·병상·ETA"),
    ]
    y = 435
    for i, (head, sub) in enumerate(labels):
        x = 100 + i * 360
        rounded(d, (x, y, x + 270, y + 170), 32, WHITE, LINE)
        d.text((x + 135, y + 58), head, font=f(34, True), fill=NAVY, anchor="mm")
        d.text((x + 135, y + 116), sub, font=f(24, False), fill=MUTED, anchor="mm")
        if i < len(labels) - 1:
            arrow(d, (x + 280, y + 85), (x + 345, y + 85), CYAN, 7)
    rounded(d, (240, 760, 1680, 885), 36, "#eafcff", CYAN)
    d.text((960, 825), "출력: 응급도 · 의심 질환 · 진료과 · 추천 병원 · 유사 사례", font=f(40, True), fill=NAVY, anchor="mm")
    footer(d, 4)
    return img


def slide_05():
    img = canvas()
    d = ImageDraw.Draw(img)
    title(d, "데이터 파이프라인", "문제정의 → 수집 → 정제 → 탐색 → 분석 → 활용")
    stages = [
        ("문제정의", "증상과 병원 사이\n판단 모델 필요"),
        ("데이터 수집", "네이버 지식인\n증상 사례"),
        ("데이터 정제", "비의료·동물·광고\n사례 제거"),
        ("데이터 탐색", "질환·증상군\n분포 확인"),
        ("분석 결과 활용", "모델·문진·추천\n서비스 연결"),
    ]
    for i, (head, sub) in enumerate(stages):
        x = 100 + i * 360
        rounded(d, (x, 300, x + 285, 595), 34, WHITE, LINE)
        pill(d, (x + 45, 335, x + 240, 385), head, CYAN if i != 4 else MINT, NAVY, 23)
        draw_text(d, (x + 142, 482), sub, 27, INK, True, anchor="mm", align="center", spacing=8)
        if i < len(stages) - 1:
            arrow(d, (x + 295, 448), (x + 350, 448), CYAN, 7)
    d.text((960, 800), "핵심: 데이터는 질환 판단용과 응급도 보조용으로 분리", font=f(42, True), fill=NAVY, anchor="mm")
    footer(d, 5)
    return img


def slide_06():
    img = canvas()
    d = ImageDraw.Draw(img)
    title(d, "데이터 이해도", "많이 모은 것보다, 같은 기준으로 연결하는 것이 중요했다")
    metrics = [
        ("3,499", "증상 사례"),
        ("149", "질환 마스터"),
        ("118", "학습 disease_id"),
        ("532", "문진 질문"),
        ("0", "unmapped row"),
    ]
    for i, (v, l) in enumerate(metrics):
        add_metric(d, (95 + i * 365, 310, 385 + i * 365, 510), v, l, MINT if i == 4 else CYAN)
    rounded(d, (245, 720, 1675, 865), 40, NAVY, NAVY)
    d.text((960, 792), "disease_id로 질환 · 질문 · 응급룰 · 병원 추천 데이터를 통합", font=f(43, True), fill=WHITE, anchor="mm")
    footer(d, 6)
    return img


def slide_07():
    img = canvas()
    d = ImageDraw.Draw(img)
    title(d, "라벨 통합: disease_master", "같은 질환이 다른 이름으로 흔들리는 문제를 줄였다")
    shadowed_card(img, (150, 300, 770, 700), 36)
    shadowed_card(img, (1150, 300, 1770, 700), 36, "#eafcff", CYAN)
    d.text((460, 370), "Before", font=f(42, True), fill=RED, anchor="mm")
    d.text((460, 505), "심근경색\n급성 심근경색\nAMI", font=f(36, True), fill=INK, anchor="mm", align="center")
    d.text((1460, 370), "After", font=f(42, True), fill=MINT, anchor="mm")
    d.text((1460, 505), "D_CAR_004\n급성 심근경색", font=f(42, True), fill=NAVY, anchor="mm", align="center")
    arrow(d, (805, 500), (1115, 500), CYAN, 10)
    d.text((960, 820), "모델·문진·응급룰이 같은 질환을 같은 ID로 바라본다", font=f(40, True), fill=NAVY, anchor="mm")
    footer(d, 7)
    return img


def slide_08():
    img = canvas()
    d = ImageDraw.Draw(img)
    title(d, "모델 설계", "자연어 증상을 의심 질환 후보로 변환")
    rounded(d, (120, 320, 520, 640), 36, WHITE, LINE)
    rounded(d, (760, 270, 1160, 690), 36, "#eafcff", CYAN)
    rounded(d, (1400, 320, 1800, 640), 36, WHITE, LINE)
    d.text((320, 395), "입력", font=f(38, True), fill=NAVY, anchor="mm")
    d.text((320, 515), "자연어\n증상 문장", font=f(34, True), fill=MUTED, anchor="mm", align="center")
    d.text((960, 370), "symptom_classifier.pkl", font=f(34, True), fill=NAVY, anchor="mm")
    d.text((960, 500), "char TF-IDF\n+ 분류 모델\n+ disease_id", font=f(32, True), fill=BLUE, anchor="mm", align="center")
    d.text((1600, 395), "출력", font=f(38, True), fill=NAVY, anchor="mm")
    d.text((1600, 515), "증상군\n진료과\n질환 TOP3", font=f(33, True), fill=MUTED, anchor="mm", align="center")
    arrow(d, (540, 480), (740, 480), CYAN, 8)
    arrow(d, (1180, 480), (1380, 480), CYAN, 8)
    d.text((960, 820), "모델 비교: baseline · Word TF-IDF · Char TF-IDF · Linear SVC", font=f(34, True), fill=NAVY, anchor="mm")
    footer(d, 8)
    return img


def slide_09():
    img = canvas()
    d = ImageDraw.Draw(img)
    title(d, "추가 문진", "질문을 많이 하는 것이 아니라, 필요한 질문만 고른다")
    rounded(d, (155, 285, 1765, 515), 44, NAVY, NAVY)
    d.text((960, 398), "질환 일치 + 증상군 일치 + 위험도 + 중요도 - 이미 말한 정보", font=f(42, True), fill=WHITE, anchor="mm")
    components = [("질환 전용", CYAN), ("위험도 반영", YELLOW), ("중복 제외", RED), ("Top 1~4", MINT)]
    for i, (txt, color) in enumerate(components):
        pill(d, (305 + i * 360, 670, 585 + i * 360, 735), txt, color, NAVY if color != RED else WHITE, 30)
    d.text((960, 875), "질문은 새로 생성하지 않고 disease_question_map 후보 중에서 랭킹", font=f(35, True), fill=NAVY, anchor="mm")
    footer(d, 9)
    return img


def slide_10():
    img = canvas()
    d = ImageDraw.Draw(img)
    title(d, "응급도와 병원 추천", "응급 데이터는 질환 판단이 아니라 위험도 계산에 사용")
    blocks = [
        ("Red flag", "호흡곤란 · 의식저하\n편마비 · 흉통", RED),
        ("Risk score", "문진 답변으로\n위험 점수 계산", YELLOW),
        ("Hospital rank", "진료과 · 병상\n거리 · ETA", CYAN),
    ]
    for i, (head, sub, color) in enumerate(blocks):
        x = 170 + i * 585
        rounded(d, (x, 315, x + 420, 655), 40, WHITE, LINE)
        pill(d, (x + 75, 365, x + 345, 425), head, color, WHITE if color == RED else NAVY, 28)
        d.text((x + 210, 535), sub, font=f(34, True), fill=INK, anchor="mm", align="center")
    rounded(d, (310, 790, 1610, 895), 34, "#eafcff", CYAN)
    d.text((960, 842), "가용 병상 0개 병원은 거리만 가까워도 우선순위 하락", font=f(36, True), fill=NAVY, anchor="mm")
    footer(d, 10)
    return img


def slide_11():
    img = canvas()
    d = ImageDraw.Draw(img)
    title(d, "실제 서비스 화면", "입력 화면과 결과 화면을 한 흐름으로 확인")
    shadowed_card(img, (85, 230, 910, 880), 36)
    paste_fit(img, ASSETS / "site_home_top.png", (115, 265, 880, 850))
    shadowed_card(img, (1010, 230, 1835, 880), 36)
    paste_fit(img, ASSETS / "site_stroke_result_top.png", (1040, 265, 1805, 850))
    pill(d, (220, 885, 775, 945), "증상 입력 · 위치 설정", CYAN, NAVY, 28)
    pill(d, (1160, 885, 1695, 945), "응급도 · 질환 · 병원 추천", MINT, NAVY, 28)
    footer(d, 11)
    return img


def slide_12():
    img = canvas()
    d = ImageDraw.Draw(img)
    title(d, "실행 검증", "정적 설명이 아니라 실제 실행으로 확인했다")
    add_metric(d, (220, 330, 560, 560), "21", "검증 항목 PASS", MINT)
    add_metric(d, (625, 330, 965, 560), "0", "FAIL", MINT)
    add_metric(d, (1030, 330, 1370, 560), "3", "주요 API 확인", CYAN)
    add_metric(d, (1435, 330, 1775, 560), "LIVE", "배포 사이트 테스트", CYAN)
    checks = "모델 로드 · disease_master · mapped CSV · backend 실행 · smoke test · /questions · /analyze"
    rounded(d, (210, 760, 1710, 875), 36, NAVY, NAVY)
    d.text((960, 817), checks, font=f(31, True), fill=WHITE, anchor="mm")
    footer(d, 12)
    return img


def slide_13():
    img = canvas(NAVY)
    d = ImageDraw.Draw(img)
    d.ellipse((-220, 650, 680, 1450), fill="#102b54")
    d.ellipse((1320, -260, 2180, 620), fill="#12345f")
    pill(d, (110, 100, 385, 158), "FINAL MESSAGE", CYAN, NAVY, 25)
    draw_text(d, (120, 275), "진단기가 아니라,\n응급도와 병원 선택을 돕는\n데이터 기반 안내 시스템", 68, WHITE, True, spacing=18)
    rounded(d, (1140, 300, 1735, 705), 42, NAVY2, "#24466d")
    bullets = ["모델 기반 질환 후보", "질환별 최소 문진", "red flag 우선", "응급/상시 병원 추천"]
    for i, b in enumerate(bullets):
        y = 390 + i * 72
        d.ellipse((1212, y - 12, 1236, y + 12), fill=CYAN)
        d.text((1260, y), b, font=f(34, True), fill="#d7f7ff", anchor="lm")
    footer(d, 13)
    return img


SLIDES = [
    slide_01,
    slide_02,
    slide_03,
    slide_04,
    slide_05,
    slide_06,
    slide_07,
    slide_08,
    slide_09,
    slide_10,
    slide_11,
    slide_12,
    slide_13,
]

NOTES = [
    "프로젝트는 강원도에서 증상 입력 후 응급도와 병원을 안내하는 시스템입니다.",
    "핵심 문제는 사용자가 증상은 말할 수 있지만, 진료과와 병원 선택은 어렵다는 점입니다.",
    "목표는 모델 예측, 최소 문진, 응급도 판단, 병원 추천 네 가지입니다.",
    "시스템은 입력부터 병원 추천까지 하나의 파이프라인으로 연결됩니다.",
    "데이터는 수집, 정제, 탐색, 분석, 활용 단계로 구성했습니다.",
    "현재 데이터는 3,499건 증상 사례와 149개 질환 마스터로 통합되어 있습니다.",
    "질환명 흔들림을 줄이기 위해 disease_id를 기준으로 통합했습니다.",
    "모델은 자연어 증상 문장을 받아 증상군, 진료과, 질환 후보 TOP3를 예측합니다.",
    "추가 문진은 고정 질문이 아니라 필요한 질문만 랭킹해서 선택합니다.",
    "응급도는 red flag와 risk score로 계산하고, 병원 추천은 병상과 ETA까지 반영합니다.",
    "실제 화면에서는 증상 입력부터 응급도와 병원 추천까지 확인할 수 있습니다.",
    "최종 검증은 21개 항목 PASS, 주요 API와 배포 사이트까지 확인했습니다.",
    "이 시스템은 진단기가 아니라 응급도와 병원 선택을 돕는 데이터 기반 안내 시스템입니다.",
]


def build_png_slides():
    VISUAL_DIR.mkdir(parents=True, exist_ok=True)
    paths = []
    for idx, maker in enumerate(SLIDES, start=1):
        img = maker()
        path = VISUAL_DIR / f"slide_{idx:02d}.png"
        img.save(path, optimize=True)
        paths.append(path)
    return paths


def build_ppt(paths):
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for path in paths:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(path), 0, 0, width=prs.slide_width, height=prs.slide_height)
    prs.save(PPTX_PATH)


def build_notes():
    lines = ["# Presentation Speaker Notes", ""]
    for idx, note in enumerate(NOTES, start=1):
        lines.append(f"## {idx}")
        lines.append(note)
        lines.append("")
    NOTES_PATH.write_text("\n".join(lines), encoding="utf-8")


def main():
    REPORTS.mkdir(exist_ok=True)
    paths = build_png_slides()
    build_ppt(paths)
    build_notes()
    print(PPTX_PATH)
    print(VISUAL_DIR)


if __name__ == "__main__":
    main()
