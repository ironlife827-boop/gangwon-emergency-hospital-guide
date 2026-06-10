from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "reports"
ASSET_DIR = OUT_DIR / "presentation_assets"
PPTX_PATH = OUT_DIR / "gangwon_emergency_hospital_guide_presentation.pptx"
SCRIPT_PATH = OUT_DIR / "presentation_speaker_notes.md"


WIDE_W = Inches(13.333)
WIDE_H = Inches(7.5)

NAVY = RGBColor(8, 20, 44)
NAVY2 = RGBColor(13, 32, 63)
CYAN = RGBColor(0, 203, 232)
CYAN_DARK = RGBColor(0, 150, 180)
WHITE = RGBColor(255, 255, 255)
MUTED = RGBColor(116, 134, 160)
SOFT = RGBColor(232, 246, 249)
PANEL = RGBColor(244, 248, 251)
RED = RGBColor(244, 74, 74)
YELLOW = RGBColor(255, 194, 66)
GREEN = RGBColor(35, 190, 120)
INK = RGBColor(25, 36, 55)


def add_textbox(slide, text, x, y, w, h, font_size=24, color=INK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    r = p.runs[0]
    r.font.name = "Malgun Gothic"
    r.font.size = Pt(font_size)
    r.font.color.rgb = color
    r.font.bold = bold
    return box


def add_title(slide, title, subtitle=None, dark=False):
    color = WHITE if dark else NAVY
    add_textbox(slide, title, Inches(0.72), Inches(0.46), Inches(11.9), Inches(0.55), 28, color, True)
    if subtitle:
        add_textbox(slide, subtitle, Inches(0.75), Inches(1.02), Inches(11.4), Inches(0.38), 12, MUTED if not dark else RGBColor(190, 214, 230))


def add_chip(slide, text, x, y, w, fill=CYAN, color=NAVY):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, Inches(0.32))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = fill
    shape.adjustments[0] = 0.18
    tf = shape.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    r = p.runs[0]
    r.font.name = "Malgun Gothic"
    r.font.size = Pt(10.5)
    r.font.bold = True
    r.font.color.rgb = color
    return shape


def add_card(slide, x, y, w, h, fill=PANEL, line=RGBColor(218, 230, 238), radius=True):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def add_bullet_list(slide, items, x, y, w, h, font_size=17, color=INK, bullet_color=CYAN):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.margin_top = Inches(0.02)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"●  {item}"
        p.space_after = Pt(9)
        p.level = 0
        r = p.runs[0]
        r.font.name = "Malgun Gothic"
        r.font.size = Pt(font_size)
        r.font.color.rgb = color
    return box


def add_arrow(slide, x1, y1, x2, y2, color=MUTED, width=2):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    line.line.color.rgb = color
    line.line.width = Pt(width)
    line.line.end_arrowhead = True
    return line


def add_step(slide, num, title, desc, x, y, w=Inches(2.1), fill=WHITE):
    add_card(slide, x, y, w, Inches(1.18), fill=fill, line=RGBColor(214, 229, 238))
    circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, x + Inches(0.16), y + Inches(0.18), Inches(0.32), Inches(0.32))
    circle.fill.solid()
    circle.fill.fore_color.rgb = CYAN
    circle.line.color.rgb = CYAN
    tf = circle.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = str(num)
    p.alignment = PP_ALIGN.CENTER
    r = p.runs[0]
    r.font.name = "Malgun Gothic"
    r.font.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = NAVY
    add_textbox(slide, title, x + Inches(0.58), y + Inches(0.15), w - Inches(0.7), Inches(0.28), 13.5, NAVY, True)
    add_textbox(slide, desc, x + Inches(0.18), y + Inches(0.62), w - Inches(0.3), Inches(0.4), 10.5, MUTED)


def add_big_number(slide, number, label, x, y, w, accent=CYAN):
    add_card(slide, x, y, w, Inches(1.15), fill=WHITE, line=RGBColor(219, 232, 240))
    add_textbox(slide, number, x + Inches(0.16), y + Inches(0.18), w - Inches(0.32), Inches(0.42), 27, accent, True, PP_ALIGN.CENTER)
    add_textbox(slide, label, x + Inches(0.12), y + Inches(0.72), w - Inches(0.24), Inches(0.26), 10.5, MUTED, False, PP_ALIGN.CENTER)


def add_footer(slide, idx):
    add_textbox(slide, f"{idx:02d}", Inches(12.45), Inches(7.03), Inches(0.5), Inches(0.24), 9, MUTED, True, PP_ALIGN.RIGHT)


def add_full_bg(slide, color=WHITE):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def fit_picture(slide, path, x, y, w, h):
    pic = slide.shapes.add_picture(str(path), x, y, width=w)
    if pic.height > h:
        pic.height = h
    pic.left = x + int((w - pic.width) / 2)
    pic.top = y + int((h - pic.height) / 2)
    return pic


def build_deck():
    prs = Presentation()
    prs.slide_width = WIDE_W
    prs.slide_height = WIDE_H
    blank = prs.slide_layouts[6]

    slides = []

    # 1 Title
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, NAVY)
    add_textbox(slide, "Gangwon Emergency\nHospital Guide", Inches(0.72), Inches(1.15), Inches(8.4), Inches(1.55), 43, WHITE, True)
    add_textbox(slide, "강원도 맞춤형 응급 및 상시 병원 안내 시스템", Inches(0.78), Inches(2.92), Inches(7.8), Inches(0.35), 17, RGBColor(198, 232, 241), False)
    add_bullet_list(slide, ["자연어 증상 입력", "모델 기반 의심 질환 TOP3", "응급도 + 병원 추천"], Inches(0.82), Inches(3.72), Inches(4.6), Inches(1.4), 18, WHITE)
    add_card(slide, Inches(7.25), Inches(1.18), Inches(5.15), Inches(4.45), fill=NAVY2, line=RGBColor(34, 68, 105))
    fit_picture(slide, ASSET_DIR / "site_home_top.png", Inches(7.48), Inches(1.42), Inches(4.7), Inches(3.9))
    add_chip(slide, "DATA MINING PROJECT", Inches(0.78), Inches(0.68), Inches(1.75), fill=CYAN)
    add_footer(slide, 1)
    slides.append(("도입", "이 프로젝트는 단순 병원 검색이 아니라 증상과 병원 사이의 판단 과정을 데이터와 모델로 연결한 시스템입니다."))

    # 2 Problem
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, WHITE)
    add_title(slide, "문제 정의", "사용자는 증상을 말하지만, 어떤 진료과와 병원을 가야 할지 모른다")
    add_card(slide, Inches(0.95), Inches(2.0), Inches(3.2), Inches(2.2), WHITE)
    add_textbox(slide, "사용자", Inches(1.2), Inches(2.35), Inches(2.7), Inches(0.4), 25, NAVY, True, PP_ALIGN.CENTER)
    add_textbox(slide, "“갑자기 왼팔에 힘이 없고\n말이 어눌해요”", Inches(1.15), Inches(3.0), Inches(2.8), Inches(0.72), 15, MUTED, False, PP_ALIGN.CENTER)
    add_card(slide, Inches(5.05), Inches(1.75), Inches(3.25), Inches(2.7), fill=SOFT, line=CYAN)
    add_textbox(slide, "비어 있던 중간 판단", Inches(5.35), Inches(2.1), Inches(2.65), Inches(0.35), 20, NAVY, True, PP_ALIGN.CENTER)
    add_bullet_list(slide, ["증상군", "진료과", "의심 질환", "응급도"], Inches(5.55), Inches(2.75), Inches(2.2), Inches(1.0), 14, INK)
    add_card(slide, Inches(9.2), Inches(2.0), Inches(3.2), Inches(2.2), WHITE)
    add_textbox(slide, "결과", Inches(9.47), Inches(2.35), Inches(2.65), Inches(0.4), 25, NAVY, True, PP_ALIGN.CENTER)
    add_textbox(slide, "응급도와\n가야 할 병원", Inches(9.55), Inches(3.02), Inches(2.45), Inches(0.7), 17, MUTED, False, PP_ALIGN.CENTER)
    add_arrow(slide, Inches(4.15), Inches(3.1), Inches(5.05), Inches(3.1), CYAN_DARK, 3)
    add_arrow(slide, Inches(8.3), Inches(3.1), Inches(9.2), Inches(3.1), CYAN_DARK, 3)
    add_textbox(slide, "목표: 자연어 증상과 병원 추천 사이의 판단 과정을 데이터마이닝으로 구현", Inches(1.1), Inches(5.55), Inches(11.2), Inches(0.4), 22, NAVY, True, PP_ALIGN.CENTER)
    add_footer(slide, 2)
    slides.append(("문제 정의", "사용자가 증상을 입력하면 바로 병원을 검색하는 것이 아니라, 중간에 증상군과 진료과와 응급도를 판단하는 모델이 필요했습니다."))

    # 3 Goals
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, WHITE)
    add_title(slide, "프로젝트 목표", "많이 묻지 않고, 위험한 증상은 놓치지 않고, 가까운 병원으로 연결한다")
    goals = [
        ("01", "모델 기반 예측", "네이버 지식인 증상 데이터로\n증상군·진료과·질환 후보 예측"),
        ("02", "최소 문진", "질환별 질문 후보 중\n필요한 1~4개만 선택"),
        ("03", "응급도 판단", "red flag와 risk score로\n위험도 보정"),
        ("04", "병원 추천", "응급/상시 병원, 병상,\n거리·ETA를 함께 반영"),
    ]
    for i, (num, title, desc) in enumerate(goals):
        x = Inches(0.75 + i * 3.12)
        add_card(slide, x, Inches(2.0), Inches(2.65), Inches(2.7), fill=PANEL)
        add_textbox(slide, num, x + Inches(0.2), Inches(2.25), Inches(0.7), Inches(0.35), 19, CYAN_DARK, True)
        add_textbox(slide, title, x + Inches(0.2), Inches(2.82), Inches(2.2), Inches(0.35), 20, NAVY, True)
        add_textbox(slide, desc, x + Inches(0.22), Inches(3.45), Inches(2.15), Inches(0.7), 13, MUTED)
    add_textbox(slide, "핵심 메시지", Inches(1.0), Inches(5.65), Inches(2.0), Inches(0.32), 14, CYAN_DARK, True)
    add_textbox(slide, "질환 판단은 모델이 중심, 응급도 데이터는 위험도 계산 보조로 분리했습니다.", Inches(2.55), Inches(5.6), Inches(9.3), Inches(0.42), 21, NAVY, True)
    add_footer(slide, 3)
    slides.append(("목표", "우리는 모델 예측, 최소 문진, 응급도 보정, 병원 추천이라는 네 가지 목표를 잡았습니다."))

    # 4 Architecture
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, WHITE)
    add_title(slide, "전체 구조", "입력부터 결과까지 하나의 파이프라인으로 연결")
    steps = [
        ("사용자 입력", "자연어 증상"),
        ("증상 모델", "TOP3 질환 후보"),
        ("문진 선택", "질환별 질문 랭킹"),
        ("응급도 계산", "red flag + risk"),
        ("병원 추천", "진료과·병상·ETA"),
    ]
    for i, (title, desc) in enumerate(steps):
        x = Inches(0.65 + i * 2.55)
        add_step(slide, i + 1, title, desc, x, Inches(2.25), Inches(2.05))
        if i < len(steps) - 1:
            add_arrow(slide, x + Inches(2.05), Inches(2.84), x + Inches(2.48), Inches(2.84), CYAN_DARK, 2.5)
    add_card(slide, Inches(1.05), Inches(4.72), Inches(11.2), Inches(1.22), fill=SOFT, line=CYAN)
    add_textbox(slide, "출력", Inches(1.35), Inches(5.02), Inches(0.85), Inches(0.28), 14, CYAN_DARK, True)
    add_textbox(slide, "응급도 · 의심 질환 · 추천 진료과 · 추천 병원 · 유사 사례", Inches(2.12), Inches(4.93), Inches(9.5), Inches(0.44), 24, NAVY, True)
    add_footer(slide, 4)
    slides.append(("전체 구조", "사용자 입력부터 모델, 문진, 응급도 계산, 병원 추천까지 한 흐름으로 이어집니다."))

    # 5 Data flow
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, WHITE)
    add_title(slide, "데이터 수집과 정제", "사용자 표현 데이터와 응급 판단 데이터를 역할별로 분리")
    lanes = [
        ("수집", "네이버 지식인\n증상 사례"),
        ("정제", "비의료·동물·광고\n사례 제거"),
        ("라벨링", "증상군·진료과\n의심 질환"),
        ("통합", "disease_id\n매핑"),
        ("활용", "모델 학습\n문진·응급도"),
    ]
    for i, (title, desc) in enumerate(lanes):
        x = Inches(0.85 + i * 2.45)
        add_card(slide, x, Inches(2.08), Inches(1.95), Inches(2.15), fill=WHITE)
        add_textbox(slide, title, x + Inches(0.15), Inches(2.42), Inches(1.65), Inches(0.34), 19, NAVY, True, PP_ALIGN.CENTER)
        add_textbox(slide, desc, x + Inches(0.15), Inches(3.08), Inches(1.65), Inches(0.7), 13, MUTED, False, PP_ALIGN.CENTER)
        if i < len(lanes) - 1:
            add_arrow(slide, x + Inches(1.95), Inches(3.15), x + Inches(2.38), Inches(3.15), CYAN_DARK, 2)
    add_bullet_list(slide, ["질환 판단용 데이터와 응급도 보조 데이터를 분리", "질환명 흔들림은 disease_master와 disease_id로 통합", "mapped CSV unmapped row 0개"], Inches(1.15), Inches(5.1), Inches(10.8), Inches(1.1), 17)
    add_footer(slide, 5)
    slides.append(("데이터", "데이터는 수집, 정제, 라벨링, 통합을 거쳤고 질환 판단과 응급도 계산의 역할을 분리했습니다."))

    # 6 Data stats
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, WHITE)
    add_title(slide, "데이터 규모", "평가자가 바로 확인할 수 있는 핵심 숫자")
    stats = [
        ("3,499", "네이버 증상 사례"),
        ("149", "disease master 질환"),
        ("118", "학습 대상 disease_id"),
        ("532", "질환별 문진 질문"),
        ("0", "unmapped rows"),
    ]
    for i, (num, label) in enumerate(stats):
        x = Inches(0.65 + i * 2.52)
        add_big_number(slide, num, label, x, Inches(2.25), Inches(2.05), CYAN_DARK if i != 4 else GREEN)
    add_card(slide, Inches(1.05), Inches(4.65), Inches(11.2), Inches(1.25), fill=NAVY, line=NAVY)
    add_textbox(slide, "데이터 이해도 포인트", Inches(1.35), Inches(4.96), Inches(2.4), Inches(0.32), 14, CYAN, True)
    add_textbox(slide, "단순히 많이 모은 것이 아니라, 질환·질문·응급룰·병원 데이터를 같은 disease_id 체계로 연결했습니다.", Inches(3.25), Inches(4.84), Inches(8.5), Inches(0.58), 20, WHITE, True)
    add_footer(slide, 6)
    slides.append(("데이터 규모", "현재 3,499건 증상 사례와 149개 질환 마스터를 disease_id로 연결했고, unmapped row는 0개입니다."))

    # 7 Model
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, WHITE)
    add_title(slide, "모델 설계", "자연어 증상을 의심 질환 TOP3 후보로 변환")
    add_card(slide, Inches(0.9), Inches(2.0), Inches(3.15), Inches(2.55), fill=SOFT, line=CYAN)
    add_textbox(slide, "입력", Inches(1.22), Inches(2.32), Inches(2.55), Inches(0.35), 22, NAVY, True, PP_ALIGN.CENTER)
    add_textbox(slide, "“열이 나고 기침이 있으며\n냄새를 잘 못 맡겠습니다”", Inches(1.12), Inches(3.05), Inches(2.75), Inches(0.7), 14.5, MUTED, False, PP_ALIGN.CENTER)
    add_card(slide, Inches(5.1), Inches(1.8), Inches(3.1), Inches(2.95), fill=WHITE)
    add_textbox(slide, "symptom_classifier.pkl", Inches(5.35), Inches(2.18), Inches(2.55), Inches(0.38), 18, NAVY, True, PP_ALIGN.CENTER)
    add_bullet_list(slide, ["char TF-IDF", "Linear SVC 계열", "disease_id 모델 추가"], Inches(5.45), Inches(2.82), Inches(2.3), Inches(1.0), 13.5, INK)
    add_card(slide, Inches(9.18), Inches(2.0), Inches(3.25), Inches(2.55), fill=SOFT, line=CYAN)
    add_textbox(slide, "출력", Inches(9.45), Inches(2.32), Inches(2.75), Inches(0.35), 22, NAVY, True, PP_ALIGN.CENTER)
    add_textbox(slide, "증상군 · 진료과\n의심 질환 TOP3", Inches(9.42), Inches(3.06), Inches(2.75), Inches(0.7), 17, MUTED, False, PP_ALIGN.CENTER)
    add_arrow(slide, Inches(4.05), Inches(3.3), Inches(5.1), Inches(3.3), CYAN_DARK, 3)
    add_arrow(slide, Inches(8.2), Inches(3.3), Inches(9.18), Inches(3.3), CYAN_DARK, 3)
    add_textbox(slide, "모델 비교 근거: Majority baseline, Word TF-IDF, Char TF-IDF, Linear SVC 비교", Inches(1.1), Inches(5.65), Inches(11.0), Inches(0.35), 17, MUTED, False, PP_ALIGN.CENTER)
    add_footer(slide, 7)
    slides.append(("모델", "모델은 사용자의 자연어 문장을 받아 증상군과 진료과, disease_id 기반 질환 후보 TOP3를 예측합니다."))

    # 8 Question
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, WHITE)
    add_title(slide, "추가 문진 로직", "고정 질문이 아니라, 필요한 질문만 고른다")
    add_card(slide, Inches(0.95), Inches(1.92), Inches(5.3), Inches(3.8), fill=PANEL)
    add_textbox(slide, "질문 점수", Inches(1.3), Inches(2.25), Inches(1.6), Inches(0.35), 21, NAVY, True)
    formula = "질환 일치 + 증상군 일치 + risk_score + importance - 이미 입력한 정보"
    add_textbox(slide, formula, Inches(1.35), Inches(3.05), Inches(4.45), Inches(0.85), 24, CYAN_DARK, True, PP_ALIGN.CENTER)
    add_textbox(slide, "Top 1~4개만 반환", Inches(1.55), Inches(4.58), Inches(4.05), Inches(0.38), 19, NAVY, True, PP_ALIGN.CENTER)
    add_card(slide, Inches(7.0), Inches(1.92), Inches(5.05), Inches(3.8), fill=WHITE)
    add_textbox(slide, "예시", Inches(7.35), Inches(2.25), Inches(1.1), Inches(0.35), 21, NAVY, True)
    add_bullet_list(slide, ["뇌졸중 예측 -> 얼굴 처짐, 심한 두통 질문", "감전 예측 -> 전기손상 전용 질문", "감기 예측 -> 과도한 응급 질문 제외"], Inches(7.35), Inches(3.0), Inches(4.25), Inches(1.6), 16)
    add_footer(slide, 8)
    slides.append(("문진", "추가 질문은 고정 4문항이 아니라, 질환과 증상군, 위험도, 중복 여부를 계산해 필요한 질문만 고릅니다."))

    # 9 Triage & recommendation
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, WHITE)
    add_title(slide, "응급도와 병원 추천", "위험한 증상은 응급실, 일반 증상은 상시 병원으로")
    cols = [
        ("Red flag", ["호흡곤란", "의식저하", "편마비", "흉통"], RED),
        ("Risk score", ["문진 답변", "위험도 점수", "10점 clamp"], YELLOW),
        ("Hospital rank", ["진료과", "가용 병상", "거리·ETA"], CYAN_DARK),
    ]
    for i, (title, items, color) in enumerate(cols):
        x = Inches(0.85 + i * 4.15)
        add_card(slide, x, Inches(2.05), Inches(3.55), Inches(3.15), fill=WHITE)
        add_chip(slide, title, x + Inches(0.28), Inches(2.38), Inches(1.5), fill=color, color=WHITE if color != YELLOW else NAVY)
        add_bullet_list(slide, items, x + Inches(0.38), Inches(3.18), Inches(2.8), Inches(1.2), 16)
    add_textbox(slide, "병상 0개 응급병원은 거리만 가까워도 우선순위가 내려갑니다.", Inches(1.2), Inches(5.85), Inches(10.85), Inches(0.36), 21, NAVY, True, PP_ALIGN.CENTER)
    add_footer(slide, 9)
    slides.append(("응급도/추천", "red flag는 모델보다 우선하고, 병원 추천은 진료과와 병상, ETA를 함께 반영합니다."))

    # 10 UI home
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, WHITE)
    add_title(slide, "실제 화면 1: 증상 입력", "사용자는 증상을 자연어로 입력하고 위치를 설정한다")
    add_card(slide, Inches(0.75), Inches(1.55), Inches(11.85), Inches(5.35), fill=WHITE)
    fit_picture(slide, ASSET_DIR / "site_home_top.png", Inches(1.0), Inches(1.78), Inches(11.35), Inches(4.9))
    add_footer(slide, 10)
    slides.append(("화면 1", "사용자는 긴 설명 없이 증상과 위치만 입력하면 문진을 시작할 수 있습니다."))

    # 11 UI result
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, WHITE)
    add_title(slide, "실제 화면 2: 결과", "응급도, 의심 질환, 추천 병원이 한 화면에 정리된다")
    add_card(slide, Inches(0.75), Inches(1.38), Inches(7.25), Inches(5.65), fill=WHITE)
    fit_picture(slide, ASSET_DIR / "site_stroke_result_top.png", Inches(0.95), Inches(1.58), Inches(6.85), Inches(5.2))
    add_card(slide, Inches(8.35), Inches(1.75), Inches(3.85), Inches(1.1), fill=SOFT, line=CYAN)
    add_textbox(slide, "뇌졸중 예시", Inches(8.62), Inches(2.05), Inches(3.25), Inches(0.34), 22, NAVY, True, PP_ALIGN.CENTER)
    add_bullet_list(slide, ["응급도: 긴급", "질환: 뇌졸중", "진료과: 신경과", "병원: 응급의료기관"], Inches(8.55), Inches(3.25), Inches(3.45), Inches(1.6), 17)
    add_footer(slide, 11)
    slides.append(("화면 2", "최종 결과에서는 응급도, 질환, 진료과, 병원 추천이 카드 형태로 바로 보입니다."))

    # 12 Demo scenarios
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, WHITE)
    add_title(slide, "발표용 데모 시나리오", "평가자가 이해하기 쉬운 대표 입력")
    demos = [
        ("뇌졸중", "한쪽 팔 힘 빠짐 + 말 어눌"),
        ("심근경색", "가슴 짓누름 + 식은땀"),
        ("아나필락시스", "벌 쏘임 + 입술 붓고 호흡곤란"),
        ("장염", "설사 6번 + 복통"),
        ("감기", "목 아픔 + 콧물 + 낮은 열"),
    ]
    for i, (disease, text) in enumerate(demos):
        y = Inches(1.65 + i * 0.88)
        add_card(slide, Inches(1.05), y, Inches(11.1), Inches(0.64), fill=WHITE)
        add_chip(slide, disease, Inches(1.28), y + Inches(0.16), Inches(1.45), fill=CYAN if i != 2 else RED, color=NAVY if i != 2 else WHITE)
        add_textbox(slide, text, Inches(3.1), y + Inches(0.17), Inches(8.6), Inches(0.28), 16.5, NAVY, True)
    add_textbox(slide, "데모 순서 추천: 비응급 1개 -> 응급 red flag 1개 -> 문진 흐름 1개", Inches(1.1), Inches(6.35), Inches(11.1), Inches(0.35), 19, MUTED, False, PP_ALIGN.CENTER)
    add_footer(slide, 12)
    slides.append(("데모", "발표에서는 감기 같은 비응급, 뇌졸중이나 아나필락시스 같은 응급, 장염 같은 문진 흐름을 보여주면 좋습니다."))

    # 13 Validation
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, WHITE)
    add_title(slide, "실행 검증 결과", "정적 분석이 아니라 실제 실행 기준으로 확인")
    add_big_number(slide, "21", "검증 항목 PASS", Inches(1.05), Inches(2.0), Inches(2.5), GREEN)
    add_big_number(slide, "0", "FAIL", Inches(3.95), Inches(2.0), Inches(2.5), GREEN)
    add_big_number(slide, "28", "데모 입력 테스트", Inches(6.85), Inches(2.0), Inches(2.5), CYAN_DARK)
    add_big_number(slide, "TOP10", "발표 시나리오 선정", Inches(9.75), Inches(2.0), Inches(2.5), CYAN_DARK)
    add_bullet_list(slide, ["모델 로드, disease_master, mapped CSV 검증", "FastAPI backend 실제 실행", "주요 API /questions, /analyze 테스트", "배포 사이트 브라우저 테스트 완료"], Inches(1.25), Inches(4.35), Inches(10.6), Inches(1.2), 17)
    add_footer(slide, 13)
    slides.append(("검증", "최종 검증에서는 모델과 CSV, 백엔드 실행, API, 배포 사이트까지 실제로 확인했습니다."))

    # 14 Peer evaluation
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, WHITE)
    add_title(slide, "상대팀 평가 기준 대응", "채점 포인트별로 근거를 준비")
    criteria = [
        ("문제 정의", "증상 -> 응급도 -> 병원 안내"),
        ("데이터 이해", "3,499건 + disease_id 통합"),
        ("기술 난이도", "모델·문진·추천 직접 구현"),
        ("완성도", "실제 배포 + API 검증 PASS"),
    ]
    for i, (title, desc) in enumerate(criteria):
        x = Inches(0.95 + (i % 2) * 6.15)
        y = Inches(1.95 + (i // 2) * 1.75)
        add_card(slide, x, y, Inches(5.25), Inches(1.22), fill=SOFT if i % 2 == 0 else WHITE, line=RGBColor(214, 229, 238))
        add_textbox(slide, "✓", x + Inches(0.22), y + Inches(0.28), Inches(0.4), Inches(0.35), 24, GREEN, True)
        add_textbox(slide, title, x + Inches(0.78), y + Inches(0.2), Inches(1.9), Inches(0.32), 18, NAVY, True)
        add_textbox(slide, desc, x + Inches(0.78), y + Inches(0.68), Inches(3.95), Inches(0.28), 13.5, MUTED)
    add_textbox(slide, "제출 근거 문서: docs/peer_evaluation_guide.md", Inches(1.1), Inches(6.15), Inches(11.0), Inches(0.35), 20, CYAN_DARK, True, PP_ALIGN.CENTER)
    add_footer(slide, 14)
    slides.append(("평가 대응", "상대팀 평가 기준에 맞춰 문제 정의, 데이터 이해, 기술 난이도, 완성도 근거를 문서로 준비했습니다."))

    # 15 Closing
    slide = prs.slides.add_slide(blank)
    add_full_bg(slide, NAVY)
    add_textbox(slide, "최종 메시지", Inches(0.85), Inches(0.72), Inches(3.5), Inches(0.45), 24, CYAN, True)
    add_textbox(slide, "진단기가 아니라,\n응급도와 병원 선택을 돕는\n데이터 기반 안내 시스템", Inches(0.9), Inches(1.75), Inches(7.6), Inches(2.25), 38, WHITE, True)
    add_card(slide, Inches(8.2), Inches(1.65), Inches(4.2), Inches(3.45), fill=NAVY2, line=RGBColor(34, 68, 105))
    add_bullet_list(slide, ["모델 기반 질환 후보", "최소 문진", "red flag 우선", "응급/상시 병원 추천"], Inches(8.6), Inches(2.1), Inches(3.4), Inches(1.7), 17, WHITE)
    add_textbox(slide, "강원도 맞춤형 응급 및 상시 병원 안내 시스템", Inches(0.95), Inches(6.45), Inches(8.0), Inches(0.35), 16, RGBColor(190, 214, 230), False)
    add_footer(slide, 15)
    slides.append(("마무리", "핵심은 진단이 아니라, 증상에서 병원 선택까지의 의사결정 흐름을 데이터 기반으로 만든 것입니다."))

    prs.save(PPTX_PATH)

    lines = ["# Presentation Speaker Notes", ""]
    for idx, (title, note) in enumerate(slides, start=1):
        lines.append(f"## {idx}. {title}")
        lines.append("")
        lines.append(note)
        lines.append("")
    SCRIPT_PATH.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    OUT_DIR.mkdir(exist_ok=True)
    ASSET_DIR.mkdir(exist_ok=True)
    build_deck()
    print(PPTX_PATH)
    print(SCRIPT_PATH)
